# -*- coding: utf-8 -*-
"""导出：把转写文稿转成 SRT 字幕 / Word 文档

- SRT：视频剪辑软件直接可用。转写时自动生成精确时间轴版（块起止时间）；
  历史文稿从 [mm:ss] 行解析，结束时间取下一句开始（末句估 8 秒）。
- Word：标题/小节/时间戳段落/纪要，便于归档分享。

用法:
  python export.py 文稿_transcript.md          # 同目录生成 .srt + .docx
  python export.py recordings/                 # 批量转换目录下全部文稿
"""
import os
import re
import sys


def _ms(t: str) -> int:
    m, s = t.split(":")
    return (int(m) * 60 + int(s)) * 1000


def _srt_time(ms) -> str:
    ms = int(ms)  # VAD 事件路径传来的可能是 float
    h, r = divmod(ms, 3600000)
    m, r = divmod(r, 60000)
    s, ms = divmod(r, 1000)
    return f"{h:02d}:{m:02d}:{s:02d},{ms:03d}"


def parse_md(md_text: str):
    """解析文稿：返回 (小节列表, 纪要文本行)
    小节 = (标题, [(beg_ms, end_ms或None, 行文本), ...])"""
    sections, minutes = [], []
    cur_title, cur = None, []
    in_minutes = False
    for line in md_text.splitlines():
        l = line.strip()
        if l.startswith("# 会议纪要"):
            in_minutes = True
            if cur_title:
                sections.append((cur_title, cur))
                cur_title, cur = None, []
            continue
        if l.startswith("## "):
            if cur_title:
                sections.append((cur_title, cur))
            cur_title, cur = l[3:].strip(), []
            continue
        if l.startswith("# ") or l.startswith("---") or not l:
            continue
        m = re.match(r"^\[(\d{2}:\d{2})\]\s*(.*)$", l)
        if m:
            cur.append((_ms(m.group(1)), None, m.group(2)))
        elif in_minutes:
            minutes.append(l)
        elif cur_title:
            cur.append((None, None, l))  # 无时间戳行（如转写内容标题行）
    if cur_title:
        sections.append((cur_title, cur))
    return sections, minutes


def _entries_to_srt(entries):
    """[(beg_ms, end_ms或None, text), ...] → SRT 文本；缺结束时间取下一句开始/8秒"""
    out, idx = [], 1
    for i, (beg, end, text) in enumerate(entries):
        if beg is None:
            continue
        if end is None:
            nxt = entries[i + 1][0] if i + 1 < len(entries) and entries[i + 1][0] else beg + 8000
            end = min(nxt, beg + 10000)
        text = text.strip()
        if not text:
            continue
        out.append(f"{idx}\n{_srt_time(beg)} --> {_srt_time(end)}\n{text}\n")
        idx += 1
    return "\n".join(out)


def blocks_to_srt(all_blocks):
    """精确时间轴：[(beg_ms, end_ms, text), ...]（已按时间排序）→ SRT 文本"""
    return _entries_to_srt([(b, e, t) for b, e, t in all_blocks])


def md_to_srt_text(md_text: str) -> str:
    """从文稿文本生成 SRT（双轨合并按时间排序；末句估 8 秒）"""
    sections, _ = parse_md(md_text)
    entries = []
    for _, rows in sections:
        entries.extend(rows)
    entries = [e for e in sorted((e for e in entries if e[0] is not None), key=lambda e: e[0])]
    return _entries_to_srt(entries)


def md_to_docx(md_path: str, docx_path: str):
    from docx import Document
    from docx.shared import Pt, RGBColor

    md_text = open(md_path, encoding="utf-8").read()
    sections, minutes = parse_md(md_text)
    doc = Document()
    title = os.path.splitext(os.path.basename(md_path))[0].replace("_transcript", "")
    h = doc.add_heading(f"谛听转写 · {title}", level=1)
    h.runs[0].font.color.rgb = RGBColor(0x1F, 0x6F, 0xEB)

    for sec_title, rows in sections:
        if sec_title:
            doc.add_heading(sec_title, level=2)
        for beg, _, text in rows:
            if beg is None:
                continue
            p = doc.add_paragraph()
            ts = p.add_run(f"[{beg // 60000:02d}:{beg // 1000 % 60:02d}] ")
            ts.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            ts.font.size = Pt(10)
            p.add_run(text)

    if minutes:
        doc.add_page_break()
        doc.add_heading("会议纪要", level=1)
        for line in minutes:
            if re.match(r"^#{1,6} ", line):
                doc.add_heading(line.lstrip("# "), level=3)
            elif line.startswith("- [ ]"):
                doc.add_paragraph("☐ " + line[5:].strip(), style="List Bullet")
            elif line.startswith(("-", "•")):
                doc.add_paragraph(line[1:].strip(), style="List Bullet")
            else:
                doc.add_paragraph(line)
    doc.save(docx_path)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    import glob

    target = sys.argv[1]
    files = (sorted(glob.glob(os.path.join(target, "*_transcript.md")))
             if os.path.isdir(target) else [target])
    if not files:
        sys.exit(f"没有找到 *_transcript.md：{target}")
    for md in files:
        srt = os.path.splitext(md)[0] + ".srt"
        docx = os.path.splitext(md)[0] + ".docx"
        with open(md, encoding="utf-8") as f:
            open(srt, "w", encoding="utf-8").write(md_to_srt_text(f.read()))
        md_to_docx(md, docx)
        print(f"OK {md}\n   -> {srt}\n   -> {docx}")


if __name__ == "__main__":
    main()
